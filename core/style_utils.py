"""
style_utils.py — Pure Python utility functions for Excel generation.
Totally clean and independent of Streamlit.
"""

import io

def create_executive_excel(df, sheet_name="Rapport") -> bytes:
    # pandas (~1.8 s) and openpyxl (~0.8 s) are only needed while an Excel file is
    # being written, so they load here rather than at application start.
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    buffer = io.BytesIO()
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        df = pd.DataFrame([{"Information": "Aucune donnée disponible"}])

    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]
        
        # Executive Header Styling: Dark Navy background, Crisp White Bold text
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        # Apply header styling & set row height
        worksheet.row_dimensions[1].height = 28
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
            
        # Data rows formatting & auto-fit column widths
        data_font = Font(name="Calibri", size=11)
        data_alignment = Alignment(horizontal="center", vertical="center")
        
        for row_idx in range(2, len(df) + 2):
            worksheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                
        # Auto-adjust column widths cleanly + padding
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return buffer.getvalue()
