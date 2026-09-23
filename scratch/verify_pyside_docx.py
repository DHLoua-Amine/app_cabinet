import sys
from pathlib import Path
import io
import docx
from docx.oxml.ns import qn

pyside_dir = Path(r"C:\Users\amin\Desktop\zarai1_pyside")
sys.path.insert(0, str(pyside_dir / "core"))
sys.path.insert(0, str(pyside_dir))

from core.pdf_generator import generate_docx_from_transcription

b = generate_docx_from_transcription("عقد بيع", "الحمد لله في يوم السبت... نحن عبد الحميد زارعي\nالفصل الأول: باع واحال...")
out_path = pyside_dir / "scratch" / "test_pyside_out.docx"
with open(out_path, "wb") as f:
    f.write(b)

doc = docx.Document(out_path)

out_lines = []
out_lines.append("=== CHECKING PYSIDE HEADER TABLE ===")
hdr_tbl = doc.tables[0]
for row_idx, row in enumerate(hdr_tbl.rows):
    for col_idx, cell in enumerate(row.cells):
        cell_text = cell.text.strip().replace("\n", " | ")
        out_lines.append(f"Header Cell ({row_idx}, {col_idx}) text: '{cell_text}'")

out_lines.append("\n=== CHECKING PYSIDE FOOTER TABLE ===")
ftr_tbl = doc.tables[1]
for row_idx, row in enumerate(ftr_tbl.rows):
    for col_idx, cell in enumerate(row.cells):
        for p in cell.paragraphs:
            for r in p.runs:
                rPr = r._r.get_or_add_rPr()
                rFonts = rPr.find(qn('w:rFonts'))
                cs = rFonts.get(qn('w:cs')) if rFonts is not None else None
                ascii_f = rFonts.get(qn('w:ascii')) if rFonts is not None else None
                sz = r.font.size.pt if r.font.size else None
                out_lines.append(f"Footer Cell ({row_idx}, {col_idx}) text: '{r.text.replace(chr(10), ' ')}' | font: {r.font.name} | size: {sz}pt | w:cs: {cs} | w:ascii: {ascii_f}")

out_lines.append("\n=== CHECKING PYSIDE BODY PARAGRAPHS ===")
for p_idx, p in enumerate(doc.paragraphs):
    for r_idx, r in enumerate(p.runs):
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        cs = rFonts.get(qn('w:cs')) if rFonts is not None else None
        ascii_f = rFonts.get(qn('w:ascii')) if rFonts is not None else None
        sz = r.font.size.pt if r.font.size else None
        out_lines.append(f"P{p_idx} R{r_idx} text: '{r.text[:25]}' | font: {r.font.name} | size: {sz}pt | w:cs: {cs} | w:ascii: {ascii_f}")

report_file = pyside_dir / "scratch" / "pyside_docx_report.txt"
with open(report_file, "w", encoding="utf-8") as f:
    f.write("\n".join(out_lines))

print(f"Report saved to {report_file}")
